"""Parse Premier Prep League Word scoresheets for GameSheet Assistant."""
from __future__ import annotations

import re
from collections import Counter
from typing import Any, BinaryIO
from zipfile import BadZipFile, ZipFile

from docx import Document
from lxml import etree

from workflow import build_entry_steps

REGULATION_SECONDS = 25 * 60
OVERTIME_SECONDS = 5 * 60
TIME_RE = re.compile(r"^\d{1,2}:\d{2}$")
NUMBER_PLAYER_RE = re.compile(r"^#?\s*(\d+)\s*(.*)$")
WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _clean(value: object) -> str:
    return " ".join(str(value or "").replace("\u2013", "-").split()).strip()


def _clock_seconds(value: str) -> int | None:
    match = TIME_RE.fullmatch(_clean(value))
    if not match:
        return None
    minutes, seconds = map(int, match.group(0).split(":"))
    return minutes * 60 + seconds if seconds < 60 else None


def _clock_text(seconds: int) -> str:
    seconds = max(0, int(seconds))
    return f"{seconds // 60}:{seconds % 60:02d}"


def _minutes_played(value: str) -> str:
    value = _clean(value)
    if _clock_seconds(value) is not None:
        return value
    return f"{value}:00" if value.isdigit() else "0:00"


def _runs(values: list[str]) -> list[tuple[str, int, int]]:
    result: list[tuple[str, int, int]] = []
    for index, value in enumerate(values):
        if result and result[-1][0] == value:
            label, start, _ = result[-1]
            result[-1] = (label, start, index + 1)
        else:
            result.append((value, index, index + 1))
    return result


def _run_starts(values: list[str], label: str) -> list[int]:
    return [start for value, start, _ in _runs(values) if value.casefold() == label.casefold()]


def _first_run(values: list[str], label: str) -> int:
    starts = _run_starts(values, label)
    if not starts:
        raise ValueError(f"PPL scoresheet is missing the {label!r} column.")
    return starts[0]


def _header_text(file_object: BinaryIO) -> list[str]:
    file_object.seek(0)
    try:
        with ZipFile(file_object) as archive:
            names = [name for name in archive.namelist() if re.fullmatch(r"word/header\d+\.xml", name)]
            if not names:
                return []
            root = etree.fromstring(archive.read(names[0]))
            return [_clean(value) for value in root.xpath(".//w:t/text()", namespaces=WORD_NS) if _clean(value)]
    except (BadZipFile, KeyError, etree.XMLSyntaxError) as error:
        raise ValueError("The uploaded file is not a readable Word .docx scoresheet.") from error


def _header_details(values: list[str]) -> tuple[str, str, str]:
    upper = [value.upper() for value in values]
    date = venue = game_number = ""
    if "DATE" in upper:
        index = upper.index("DATE")
        if index + 1 < len(values):
            date = values[index + 1]
    if "LOCATION" in upper:
        start = upper.index("LOCATION") + 1
        end = upper.index("DATE", start) if "DATE" in upper[start:] else len(values)
        venue = " ".join(values[start:end])
    for index, value in enumerate(upper):
        if value in {"GAME #", "GAME#", "GAME"} and index + 1 < len(values):
            game_number = values[index + 1]
            break
    return date, venue, game_number


def _team_name(row: list[str], side_label: str, stop_label: str) -> str:
    meaningful: list[str] = []
    for value in row:
        if value and value not in meaningful:
            meaningful.append(value)
    excluded = {"No.", side_label, "Pos", stop_label}
    candidates = [value for value in meaningful if value not in excluded]
    if not candidates:
        raise ValueError(f"Could not read the {side_label.lower()} team name from the PPL scoresheet.")
    return candidates[0]


def _roster(rows: list[list[str]], start: int, end: int, position_column: int) -> dict[str, str]:
    result: dict[str, str] = {}
    for row in rows[start:end]:
        number = _clean(row[0])
        if not number.isdigit():
            continue
        pieces: list[str] = []
        for value in row[1:position_column]:
            value = _clean(value)
            if value and (not pieces or value != pieces[-1]):
                pieces.append(value)
        if pieces:
            result[number] = " ".join(pieces)
    return result


def _full_player(value: str, roster: dict[str, str]) -> str:
    value = _clean(value)
    if not value or value in {"--", "-"}:
        return ""
    match = NUMBER_PLAYER_RE.match(value)
    if not match:
        return value
    number, supplied_name = match.groups()
    return f"#{number} {roster.get(number, supplied_name).strip()}".strip()


def _team_for_label(label: str, home_team: str, away_team: str) -> str:
    normalized = re.sub(r"[^a-z0-9]", "", _clean(label).lower())
    matches = []
    active_period_lengths = [
        OVERTIME_SECONDS if period.startswith("OT") else REGULATION_SECONDS
        for period in active_periods
    ]
    for team in (home_team, away_team):
        words = re.findall(r"[a-z0-9]+", team.lower())
        if normalized and (normalized == "".join(words) or normalized in words or normalized == words[-1]):
            matches.append(team)
    if len(matches) == 1:
        return matches[0]
    raise ValueError(f"Could not match scoring team {label!r} to {home_team!r} or {away_team!r}.")


def _strength(value: str) -> str:
    text = _clean(value).upper()
    if text in {"", "--", "-", "EV", "ES"}:
        return "even strength"
    parts = []
    if "PP" in text:
        parts.append("power play")
    if "SH" in text:
        parts.append("short handed")
    if "EN" in text:
        parts.append("empty net")
    if "PS" in text:
        parts.append("penalty shot")
    return " / ".join(parts) or text.lower()


def _period_label(index: int) -> str:
    return ["1st", "2nd", "3rd"][index] if index < 3 else f"OT{index - 2}"


def _penalty_description(offense: str, minutes: str) -> str:
    offense = _clean(offense) or "Minor"
    canonical_offenses = {
        "cross checking": "Cross-Checking",
        "cross-checking": "Cross-Checking",
        "body checking": "Body Checking",
        "checking from behind": "Checking From Behind",
        "high stick": "High Sticking",
        "high-sticking": "High Sticking",
        "high sticking": "High Sticking",
    }
    offense = canonical_offenses.get(offense.casefold(), offense)
    length = int(minutes) if str(minutes).isdigit() else 2
    if length >= 10 or "misconduct" in offense.lower():
        kind = "Game Misconduct" if "game" in offense.lower() else "Misconduct"
    elif length == 5 or "major" in offense.lower():
        kind = "Major"
    else:
        kind = "Minor"
    base = re.sub(r"\s*-?\s*(?:minor|major|game misconduct|misconduct)\s*$", "", offense, flags=re.I).strip()
    base = base or ("Misconduct" if "Misconduct" in kind else "Body Checking")
    return f"{base} - {kind} ({length}:00)"


def _save_percentage(saves: int, goals_against: int) -> str:
    shots = saves + goals_against
    return f"{saves / shots:.3f}" if shots else "0.000"


def parse_ppl_docx(file_object: BinaryIO) -> dict[str, Any]:
    """Convert one standard PPL Word scoresheet to the shared app data shape."""
    header = _header_text(file_object)
    file_object.seek(0)
    try:
        document = Document(file_object)
    except Exception as error:
        raise ValueError("The uploaded file is not a readable Word .docx scoresheet.") from error
    if not document.tables:
        raise ValueError("No scoresheet table was found in the Word document.")
    rows = [[_clean(cell.text) for cell in row.cells] for row in document.tables[0].rows]
    if len(rows) < 55:
        raise ValueError("This Word document does not match the expected PPL scoresheet layout.")

    home_pos = _first_run(rows[0], "Pos")
    away_pos = _first_run(rows[21], "Pos")
    home_team = _team_name(rows[0], "HOME", "SCORING")
    away_team = _team_name(rows[21], "VIS", "PENALTIES")
    home_roster = _roster(rows, 1, 21, home_pos)
    away_roster = _roster(rows, 22, 42, away_pos)
    rosters = {home_team: home_roster, away_team: away_roster}

    scoring_headers = rows[1]
    per_col = _first_run(scoring_headers, "Per")
    time_col = _first_run(scoring_headers, "Time")
    team_col = _first_run(scoring_headers, "Team")
    goal_col = _first_run(scoring_headers, "Goal")
    # The two adjacent merged Assist headers collapse into one repeated run in
    # python-docx. Their grid starts are stable relative to the Goal field.
    assist_cols = [goal_col + 3, goal_col + 5]
    strength_col = _first_run(scoring_headers, "PP/SH")
    goals: list[dict[str, Any]] = []
    for row in rows[2:21]:
        period_text, remaining = row[per_col], row[time_col]
        if not period_text.isdigit() or _clock_seconds(remaining) is None:
            continue
        period_index = int(period_text) - 1
        period = _period_label(period_index)
        period_length = OVERTIME_SECONDS if period.startswith("OT") else REGULATION_SECONDS
        team = _team_for_label(row[team_col], home_team, away_team)
        roster = rosters[team]
        assists = [_full_player(row[column], roster) for column in assist_cols[:2]]
        goals.append({
            "period": period,
            "elapsed": _clock_text(period_length - (_clock_seconds(remaining) or 0)),
            "remaining": remaining,
            "team": team,
            "scorer": _full_player(row[goal_col], roster) or "Team/Unknown",
            "strength": _strength(row[strength_col]),
            "assists": [assist for assist in assists if assist],
        })

    penalty_headers = rows[22]
    penalty_per_col = _first_run(penalty_headers, "Per")
    penalty_team_col = _first_run(penalty_headers, "Team")
    player_col = _first_run(penalty_headers, "Player")
    offense_col = _first_run(penalty_headers, "Offense")
    minutes_col = _first_run(penalty_headers, "Min.")
    in_col = _first_run(penalty_headers, "In")
    out_col = _first_run(penalty_headers, "Out")
    penalties: list[dict[str, str]] = []
    for row in rows[23:42]:
        period_text, remaining = row[penalty_per_col], row[in_col]
        if not period_text.isdigit() or _clock_seconds(remaining) is None:
            continue
        period_index = int(period_text) - 1
        period = _period_label(period_index)
        period_length = OVERTIME_SECONDS if period.startswith("OT") else REGULATION_SECONDS
        team = _team_for_label(row[penalty_team_col], home_team, away_team)
        penalties.append({
            "period": period,
            "elapsed": _clock_text(period_length - (_clock_seconds(remaining) or 0)),
            "remaining": remaining,
            "time_on": row[out_col] if _clock_seconds(row[out_col]) is not None else "",
            "team": team,
            "player": _full_player(row[player_col], rosters[team]) or "Team/Bench",
            "penalty": _penalty_description(row[offense_col], row[minutes_col]),
        })

    summary_headers = rows[52]
    period_columns = []
    for label in ("1st Per", "2nd Per", "3rd Per", "O.T."):
        starts = _run_starts(summary_headers, label)
        if starts:
            period_columns.append((label, starts[0]))
    active_periods = [
        _period_label(index) for index, (_, column) in enumerate(period_columns)
        if rows[53][column].isdigit() or rows[54][column].isdigit()
    ]
    if not active_periods:
        active_periods = sorted({goal["period"] for goal in goals} | {penalty["period"] for penalty in penalties}) or ["1st", "2nd"]
    final_col = _first_run(summary_headers, "FINAL")
    home_score = int(rows[53][final_col]) if rows[53][final_col].isdigit() else sum(g["team"] == home_team for g in goals)
    away_score = int(rows[54][final_col]) if rows[54][final_col].isdigit() else sum(g["team"] == away_team for g in goals)

    saves_headers = rows[46]
    save_period_columns = [_first_run(saves_headers, label) for label in ("1st Per", "2nd Per", "3rd Per", "O.T.")[:len(active_periods)]]
    total_col = _first_run(saves_headers, "TOTAL")
    ga_col = _first_run(saves_headers, "G.A.")
    minutes_col = _first_run(saves_headers, "Minutes Played")
    goalie_period_saves = {home_team: [0] * len(active_periods), away_team: [0] * len(active_periods)}
    goalies: list[dict[str, str]] = []
    for row in rows[47:51]:
        compact_label = row[0].upper().replace(" ", "")
        team = home_team if compact_label.startswith("HOME") else away_team if compact_label.startswith("VISITOR") else ""
        number_match = re.search(r"(\d+)\s*$", row[0])
        if not team or not number_match:
            continue
        number = number_match.group(1)
        minutes = _minutes_played(row[minutes_col])
        saves_by_period = [int(row[column]) if row[column].isdigit() else 0 for column in save_period_columns]
        for index, saves in enumerate(saves_by_period):
            goalie_period_saves[team][index] += saves
        saves = int(row[total_col]) if row[total_col].isdigit() else sum(saves_by_period)
        goals_against = int(row[ga_col]) if row[ga_col].isdigit() else 0
        if minutes == "0:00" and saves == 0 and goals_against == 0:
            continue
        goalies.append({
            "team": team, "number": number, "name": rosters[team].get(number, f"Goalie {number}"),
            "minutes": minutes, "shots_against": str(saves + goals_against),
            "goals_against": str(goals_against), "saves": str(saves),
            "save_percentage": _save_percentage(saves, goals_against),
            "_period_saves": saves_by_period,
        })

    # Some PPL sheets leave Minutes Played blank even though each goalie's
    # period is explicit in the saves grid.  When one goalie has saves in only
    # one period, that row provides an unambiguous 25-minute stint and starter
    # order without relying on the missing minutes cell.
    for team in (home_team, away_team):
        team_goalies = [goalie for goalie in goalies if goalie["team"] == team]
        used_periods: set[int] = set()
        inferred: list[tuple[dict[str, Any], int]] = []
        for goalie in team_goalies:
            active = [index for index, saves in enumerate(goalie["_period_saves"]) if saves > 0]
            if len(active) == 1:
                inferred.append((goalie, active[0]))
                used_periods.add(active[0])
        if len(used_periods) == len(inferred):
            for goalie, period_index in inferred:
                if goalie["minutes"] == "0:00" and period_index < len(active_period_lengths):
                    goalie["minutes"] = _clock_text(active_period_lengths[period_index])

    for goalie in goalies:
        goalie.pop("_period_saves", None)

    goal_counts = {
        home_team: Counter(goal["period"] for goal in goals if goal["team"] == home_team),
        away_team: Counter(goal["period"] for goal in goals if goal["team"] == away_team),
    }
    home_shots = [goalie_period_saves[away_team][i] + goal_counts[home_team][period] for i, period in enumerate(active_periods)]
    away_shots = [goalie_period_saves[home_team][i] + goal_counts[away_team][period] for i, period in enumerate(active_periods)]
    shots = {
        "periods": active_periods,
        "period_lengths": [OVERTIME_SECONDS if period.startswith("OT") else REGULATION_SECONDS for period in active_periods],
        "away_team": away_team, "away": [str(value) for value in away_shots] + [str(sum(away_shots))],
        "home_team": home_team, "home": [str(value) for value in home_shots] + [str(sum(home_shots))],
    }
    date, venue, game_number = _header_details(header)
    game = {
        "date": date, "venue": venue, "status": "FINAL", "away_team": away_team,
        "away_score": str(away_score), "home_score": str(home_score), "home_team": home_team,
        "game_number": game_number, "source_type": "ppl_docx",
    }
    validation = []
    for team, score in ((away_team, away_score), (home_team, home_score)):
        parsed = sum(goal["team"] == team for goal in goals)
        validation.append({"ok": parsed == score, "label": f"{team} goals", "detail": f"Parsed {parsed}; final score {score}"})
    for goalie in goalies:
        calculated = int(goalie["saves"]) + int(goalie["goals_against"])
        validation.append({"ok": calculated == int(goalie["shots_against"]), "label": f"{goalie['team']} #{goalie['number']} {goalie['name']}", "detail": f"Saves + GA = {calculated}; shots against {goalie['shots_against']}"})
    return {
        "game": game, "shots": shots, "goals": goals, "penalties": penalties,
        "goalies": goalies, "validation": validation,
        "workflow": build_entry_steps(game, shots, goals, penalties, goalies),
        "source_url": f"PPL Word scoresheet Game {game_number}".strip(),
    }
